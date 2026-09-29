import io
import logging
from typing import Dict, List, Any
import docx
from pdfminer.high_level import extract_text as extract_pdf_text
from ..ai_agents.profile_agent import ProfileAgent

logger = logging.getLogger(__name__)

def extract_text_from_file_bytes(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from uploaded PDF, DOCX, or text file."""
    fn = filename.lower()
    try:
        if fn.endswith('.pdf'):
            text = extract_pdf_text(io.BytesIO(file_bytes))
            return text.strip()
        elif fn.endswith('.docx') or fn.endswith('.doc'):
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text:
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            return "\n".join(full_text).strip()
        else:
            # Plain text / fallback
            return file_bytes.decode('utf-8', errors='ignore').strip()
    except Exception as e:
        logger.error(f"Error extracting text from {filename}: {e}")
        return file_bytes.decode('utf-8', errors='ignore').strip()

def extract_facts_from_bytes(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Extract full structured profile and facts from uploaded document bytes."""
    text = extract_text_from_file_bytes(file_bytes, filename)
    agent = ProfileAgent()
    return agent.extract_from_text(text)

def confirm_facts(facts: List[Dict], approvals: Dict[str, bool]) -> List[Dict]:
    confirmed = []
    for f in facts:
        key = f.get("path") or f.get("fact_key")
        approved = approvals.get(key, True)
        f["verified"] = approved
        confirmed.append(f)
    return confirmed
